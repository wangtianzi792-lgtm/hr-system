#!/usr/bin/env python3
"""导入花名册 Excel 到 roster 表"""
import os
import sqlite3
import openpyxl

DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "attendance.db")
ROSTER_PATH = "/Users/jiuhua/2026年海昌花名册.xlsx"


def main():
    if not os.path.exists(ROSTER_PATH):
        print(f"找不到花名册文件: {ROSTER_PATH}")
        return

    print(f"读取花名册: {ROSTER_PATH}")
    wb = openpyxl.load_workbook(ROSTER_PATH, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if len(rows) < 2:
        print("花名册没有数据")
        return

    headers = rows[0]
    print(f"总行数: {len(rows)}, 表头列数: {len(headers)}")

    # 建立列索引映射（用精确的表头名做 key）
    col_map = {}
    for i, h in enumerate(headers):
        if h is not None:
            col_map[str(h).strip()] = i

    print(f"有效列: {len(col_map)}")
    for r in ["工号", "姓名"]:
        print(f"  {r} → 列 {col_map[r]}")

    # Excel列名 → 数据库列名的映射
    excel_to_db = {
        "工号": "工号",
        "姓名": "姓名",
        "人员类别": "人员类别",
        "事业部": "事业部",
        "部门": "部门",
        "工序": "工序",
        "岗位": "岗位",
        "职务级别": "职务级别",
        "入职时间": "入职时间",
        "试用到期日": "试用到期日",
        "联系电话": "联系电话",
        "紧急联系电话": "紧急联系电话",
        "身份证号": "身份证号",
        "性别": "性别",
        "出生日期": "出生日期",
        "年龄": "年龄",
        "民族": "民族",
        "工龄": "工龄",
        "用工形式": "用工形式",
        "体检类别": "体检类别",
        "文化程度": "文化程度",
        "学历工资": "学历工资",
        "毕业院校及专业": "毕业院校及专业",
        "户籍地址": "户籍地址",
        "扬州暂住地": "扬州暂住地",
        "合同起始": "合同起始",
        "合同终止": "合同终止",
        "签定": "签定",
        "公积金（标准）": "公积金标准",
        "工龄工资": "工龄工资",
        "档案号": "档案号",
        "厂区": "厂区",
        "序号": "序号",
    }

    db_cols = list(excel_to_db.values())  # 数据库列名顺序
    excel_cols = list(excel_to_db.keys())  # Excel列名顺序

    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    # 清空旧数据
    cursor.execute("DELETE FROM roster")
    print("已清空旧 roster 数据")

    # INSERT SQL
    cols_sql = ", ".join(db_cols)
    placeholders = ", ".join(["?"] * len(db_cols))
    insert_sql = f"INSERT INTO roster ({cols_sql}) VALUES ({placeholders})"

    inserted = 0
    skipped = 0

    for i, row in enumerate(rows[1:], start=2):
        try:
            def safe_get(excel_name):
                idx = col_map.get(excel_name)
                if idx is None:
                    return None
                v = row[idx]
                return str(v).strip() if v is not None else None

            emp_no = safe_get("工号")
            name = safe_get("姓名")

            if not emp_no or emp_no == "None" or not name or name == "None":
                skipped += 1
                continue

            if not any(c.isdigit() for c in emp_no):
                skipped += 1
                continue

            values = [safe_get(c) for c in excel_cols]
            cursor.execute(insert_sql, values)
            inserted += 1

        except Exception as e:
            emp_no_raw = row[col_map["工号"]] if col_map.get("工号") is not None else "?"
            print(f"  跳过第{i}行 (工号={emp_no_raw}): {e}")
            skipped += 1

    conn.commit()

    cursor.execute("SELECT COUNT(*) FROM roster")
    print(f"\n✅ 导入完成: {inserted} 条，跳过 {skipped} 条")

    cursor.execute("SELECT 工号, 姓名, 部门, 公积金标准 FROM roster WHERE 公积金标准 IS NOT NULL LIMIT 5")
    print("\n有公积金值的样例:")
    for r in cursor.fetchall():
        print(f"  工号={r[0]}, 姓名={r[1]}, 部门={r[2]}, 公积金={r[3]}")

    cursor.execute("SELECT COUNT(*) FROM roster WHERE 公积金标准 IS NOT NULL AND 公积金标准 != ''")
    print(f"有公积金值的人数: {cursor.fetchone()[0]}")

    conn.close()


if __name__ == "__main__":
    main()