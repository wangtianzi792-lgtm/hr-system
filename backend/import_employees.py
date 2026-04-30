#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从海昌花名册Excel导入员工数据
"""
import sys
import os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from datetime import datetime, date
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from app.core.database import Base, engine, SessionLocal
from app.models.employee import Employee
from app.models.department import Department
import openpyxl

def parse_date(value):
    """解析日期"""
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ['%Y-%m-%d', '%Y/%m/%d', '%m/%d/%Y', '%Y.%m.%d']:
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except:
                pass
    return None

def parse_int(value):
    """解析整数"""
    if not value:
        return None
    try:
        return int(float(str(value).strip()))
    except:
        return None

def get_or_create_department(db, dept_name):
    """获取或创建部门"""
    if not dept_name:
        return None
    dept = db.query(Department).filter(Department.name == dept_name).first()
    if not dept:
        dept = Department(name=dept_name, code=dept_name[:3] if len(dept_name) >= 3 else dept_name)
        db.add(dept)
        db.flush()
    return dept.id

def import_from_excel():
    """从Excel导入员工数据"""
    db = SessionLocal()
    
    try:
        # 读取Excel
        wb = openpyxl.load_workbook(r'C:\Users\Administrator\Desktop\2026年海昌花名册.xlsx')
        
        # 导入正式工
        ws = wb.worksheets[1]  # 正式工sheet
        print(f"正式工行数: {ws.max_row}")
        
        # 获取表头
        headers = {}
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col)
            if cell.value:
                headers[str(cell.value).strip()] = col
        
        print("表头:", list(headers.keys())[:10])
        
        # 清空现有员工数据（先清空关联表）
        from app.models.attendance import AttendanceRecord
        from app.models.scheduling import ShiftAssignment, ScheduleRecord
        
        db.query(AttendanceRecord).delete()
        db.query(ShiftAssignment).delete()
        db.query(ScheduleRecord).delete()
        db.query(Employee).delete()
        db.commit()
        print("已清空现有员工数据")
        
        count = 0
        for row in range(2, ws.max_row + 1):
            try:
                # 获取工号，如果没有则跳过
                emp_no = ws.cell(row=row, column=headers.get('工号', 4)).value
                if not emp_no:
                    continue
                
                emp_no = str(emp_no).strip()
                name = ws.cell(row=row, column=headers.get('姓名', 5)).value
                if not name:
                    continue
                
                # 部门处理
                dept_name = ws.cell(row=row, column=headers.get('部门', 8)).value
                department_id = get_or_create_department(db, dept_name) if dept_name else None
                
                # 创建员工
                employee = Employee(
                    employee_no=emp_no,
                    name=str(name).strip(),
                    archive_no=str(ws.cell(row=row, column=headers.get('档案号', 2)).value).strip() if ws.cell(row=row, column=headers.get('档案号', 2)).value else None,
                    factory=str(ws.cell(row=row, column=headers.get('厂区', 3)).value).strip() if ws.cell(row=row, column=headers.get('厂区', 3)).value else None,
                    person_type=str(ws.cell(row=row, column=headers.get('人员类别', 6)).value).strip() if ws.cell(row=row, column=headers.get('人员类别', 6)).value else None,
                    business_unit=str(ws.cell(row=row, column=headers.get('事业部', 7)).value).strip() if ws.cell(row=row, column=headers.get('事业部', 7)).value else None,
                    department_id=department_id,
                    process=str(ws.cell(row=row, column=headers.get('工序', 9)).value).strip() if ws.cell(row=row, column=headers.get('工序', 9)).value else None,
                    position=str(ws.cell(row=row, column=headers.get('岗位', 10)).value).strip() if ws.cell(row=row, column=headers.get('岗位', 10)).value else None,
                    medical_exam_type=str(ws.cell(row=row, column=headers.get('体检—工种', 11)).value).strip() if ws.cell(row=row, column=headers.get('体检—工种', 11)).value else None,
                    medical_exam_factor=str(ws.cell(row=row, column=headers.get('体检—有害因素', 12)).value).strip() if ws.cell(row=row, column=headers.get('体检—有害因素', 12)).value else None,
                    dept_audit=str(ws.cell(row=row, column=headers.get('部门—审计口径', 13)).value).strip() if ws.cell(row=row, column=headers.get('部门—审计口径', 13)).value else None,
                    position_audit=str(ws.cell(row=row, column=headers.get('岗位—审计口径', 14)).value).strip() if ws.cell(row=row, column=headers.get('岗位—审计口径', 14)).value else None,
                    job_level=str(ws.cell(row=row, column=headers.get('职务级别', 15)).value).strip() if ws.cell(row=row, column=headers.get('职务级别', 15)).value else None,
                    entry_date=parse_date(ws.cell(row=row, column=headers.get('入职时间', 16)).value),
                    probation_end_date=parse_date(ws.cell(row=row, column=headers.get('试用到期日', 17)).value),
                    phone=str(ws.cell(row=row, column=headers.get('联系电话', 18)).value).strip() if ws.cell(row=row, column=headers.get('联系电话', 18)).value else None,
                    emergency_phone=str(ws.cell(row=row, column=headers.get('紧急联系电话', 19)).value).strip() if ws.cell(row=row, column=headers.get('紧急联系电话', 19)).value else None,
                    id_card=str(ws.cell(row=row, column=headers.get('身份证号', 20)).value).strip() if ws.cell(row=row, column=headers.get('身份证号', 20)).value else None,
                    gender=str(ws.cell(row=row, column=headers.get('性别', 21)).value).strip() if ws.cell(row=row, column=headers.get('性别', 21)).value else None,
                    birth_date=parse_date(ws.cell(row=row, column=headers.get('出生日期', 22)).value),
                    age=parse_int(ws.cell(row=row, column=headers.get('年龄', 23)).value),
                    ethnicity=str(ws.cell(row=row, column=headers.get('民族', 24)).value).strip() if ws.cell(row=row, column=headers.get('民族', 24)).value else None,
                    work_years=str(ws.cell(row=row, column=headers.get('工龄', 25)).value).strip() if ws.cell(row=row, column=headers.get('工龄', 25)).value else None,
                    employment_type=str(ws.cell(row=row, column=headers.get('用工形式', 26)).value).strip() if ws.cell(row=row, column=headers.get('用工形式', 26)).value else None,
                    medical_category=str(ws.cell(row=row, column=headers.get('体检类别', 27)).value).strip() if ws.cell(row=row, column=headers.get('体检类别', 27)).value else None,
                    employee_group=str(ws.cell(row=row, column=headers.get('员工', 28)).value).strip() if ws.cell(row=row, column=headers.get('员工', 28)).value else None,
                    education=str(ws.cell(row=row, column=headers.get('文化程度', 29)).value).strip() if ws.cell(row=row, column=headers.get('文化程度', 29)).value else None,
                    education_salary=str(ws.cell(row=row, column=headers.get('学历工资', 30)).value).strip() if ws.cell(row=row, column=headers.get('学历工资', 30)).value else None,
                    school_major=str(ws.cell(row=row, column=headers.get('毕业院校及专业', 31)).value).strip() if ws.cell(row=row, column=headers.get('毕业院校及专业', 31)).value else None,
                    household_address=str(ws.cell(row=row, column=headers.get('户籍地址', 32)).value).strip() if ws.cell(row=row, column=headers.get('户籍地址', 32)).value else None,
                    temporary_address=str(ws.cell(row=row, column=headers.get('扬州暂住地', 33)).value).strip() if ws.cell(row=row, column=headers.get('扬州暂住地', 33)).value else None,
                    contract_start=parse_date(ws.cell(row=row, column=headers.get('合同起始', 34)).value),
                    contract_end=parse_date(ws.cell(row=row, column=headers.get('合同终止', 35)).value),
                    contract_signed=str(ws.cell(row=row, column=headers.get('签定', 36)).value).strip() if ws.cell(row=row, column=headers.get('签定', 36)).value else None,
                    provident_fund=str(ws.cell(row=row, column=headers.get('公积金（标准）', 37)).value).strip() if ws.cell(row=row, column=headers.get('公积金（标准）', 37)).value else None,
                    retirement_date=parse_date(ws.cell(row=row, column=headers.get('退休人员日期', 38)).value),
                    work_years_salary=str(ws.cell(row=row, column=headers.get('工龄工资', 40)).value).strip() if ws.cell(row=row, column=headers.get('工龄工资', 40)).value else None,
                    training=str(ws.cell(row=row, column=headers.get('参加培训', 41)).value).strip() if ws.cell(row=row, column=headers.get('参加培训', 41)).value else None,
                    status='active'
                )
                
                db.add(employee)
                count += 1
                
                if count % 50 == 0:
                    db.commit()
                    print(f"已导入 {count} 人...")
                    
            except Exception as e:
                print(f"第{row}行导入失败: {e}")
                continue
        
        db.commit()
        print(f"正式工导入完成，共 {count} 人")
        
        # 导入劳务工
        ws2 = wb.worksheets[2]  # 劳务工sheet
        print(f"\n劳务工行数: {ws2.max_row}")
        
        headers2 = {}
        for col in range(1, ws2.max_column + 1):
            cell = ws2.cell(row=1, column=col)
            if cell.value:
                headers2[str(cell.value).strip()] = col
        
        count2 = 0
        for row in range(2, ws2.max_row + 1):
            try:
                emp_no = ws2.cell(row=row, column=headers2.get('工号', 4)).value
                if not emp_no:
                    continue
                
                emp_no = str(emp_no).strip()
                name = ws2.cell(row=row, column=headers2.get('姓名', 5)).value
                if not name:
                    continue
                
                dept_name = ws2.cell(row=row, column=headers2.get('部门', 8)).value
                department_id = get_or_create_department(db, dept_name) if dept_name else None
                
                employee = Employee(
                    employee_no=emp_no,
                    name=str(name).strip(),
                    archive_no=str(ws2.cell(row=row, column=headers2.get('档案号', 2)).value).strip() if ws2.cell(row=row, column=headers2.get('档案号', 2)).value else None,
                    factory=str(ws2.cell(row=row, column=headers2.get('厂区', 3)).value).strip() if ws2.cell(row=row, column=headers2.get('厂区', 3)).value else None,
                    person_type='劳务工',
                    business_unit=str(ws2.cell(row=row, column=headers2.get('事业部', 7)).value).strip() if ws2.cell(row=row, column=headers2.get('事业部', 7)).value else None,
                    department_id=department_id,
                    process=str(ws2.cell(row=row, column=headers2.get('工序', 9)).value).strip() if ws2.cell(row=row, column=headers2.get('工序', 9)).value else None,
                    position=str(ws2.cell(row=row, column=headers2.get('岗位', 10)).value).strip() if ws2.cell(row=row, column=headers2.get('岗位', 10)).value else None,
                    entry_date=parse_date(ws2.cell(row=row, column=headers2.get('入职时间', 16)).value),
                    phone=str(ws2.cell(row=row, column=headers2.get('联系电话', 18)).value).strip() if ws2.cell(row=row, column=headers2.get('联系电话', 18)).value else None,
                    id_card=str(ws2.cell(row=row, column=headers2.get('身份证号', 20)).value).strip() if ws2.cell(row=row, column=headers2.get('身份证号', 20)).value else None,
                    gender=str(ws2.cell(row=row, column=headers2.get('性别', 21)).value).strip() if ws2.cell(row=row, column=headers2.get('性别', 21)).value else None,
                    birth_date=parse_date(ws2.cell(row=row, column=headers2.get('出生日期', 22)).value),
                    age=parse_int(ws2.cell(row=row, column=headers2.get('年龄', 23)).value),
                    ethnicity=str(ws2.cell(row=row, column=headers2.get('民族', 24)).value).strip() if ws2.cell(row=row, column=headers2.get('民族', 24)).value else None,
                    employment_type=str(ws2.cell(row=row, column=headers2.get('用工形式', 26)).value).strip() if ws2.cell(row=row, column=headers2.get('用工形式', 26)).value else None,
                    education=str(ws2.cell(row=row, column=headers2.get('文化程度', 28)).value).strip() if ws2.cell(row=row, column=headers2.get('文化程度', 28)).value else None,
                    contract_start=parse_date(ws2.cell(row=row, column=headers2.get('合同起始', 33)).value),
                    contract_end=parse_date(ws2.cell(row=row, column=headers2.get('合同终止', 34)).value),
                    status='active'
                )
                
                db.add(employee)
                count2 += 1
                
                if count2 % 50 == 0:
                    db.commit()
                    print(f"已导入 {count2} 人...")
                    
            except Exception as e:
                print(f"劳务工第{row}行导入失败: {e}")
                continue
        
        db.commit()
        print(f"劳务工导入完成，共 {count2} 人")
        print(f"\n总计导入: {count + count2} 人")
        
    except Exception as e:
        print(f"导入失败: {e}")
        import traceback
        traceback.print_exc()
    finally:
        db.close()

if __name__ == '__main__':
    # 创建表
    Base.metadata.create_all(bind=engine)
    print("数据库表已创建/更新")
    
    # 导入数据
    import_from_excel()
